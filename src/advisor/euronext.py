"""Import of the PEA PME list published by Euronext (xlsx file dropped in by hand)."""

from __future__ import annotations

import re
import sqlite3
from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook

ISIN_RE = re.compile(r"^[A-Z]{2}[A-Z0-9]{9}[0-9]$")
DATE_RE = re.compile(r"(\d{2})/(\d{2})/(\d{4})")

# The file has bilingual headers ("Société/Company"); each column is found by its English keyword.
COLUMNS = {
    "name": "company",
    "isin": "isin",
    "market": "market",
    "segment": "compartment",
    "country": "country",
}


@dataclass
class Company:
    isin: str
    name: str
    market: str | None
    segment: str | None
    country: str | None


@dataclass
class EuronextList:
    file: Path
    list_date: date | None
    companies: list[Company]
    rejected_rows: list[tuple[int, str]] = field(default_factory=list)


@dataclass
class Changes:
    added: list[Company]
    removed: list[sqlite3.Row]
    returned: list[Company]   # removed by an earlier list, present again
    unchanged: int


def is_valid_isin(isin: str) -> bool:
    """Format and check digit (Luhn over the digits obtained by converting letters)."""
    if not ISIN_RE.match(isin):
        return False
    digits = "".join(str(int(c, 36)) for c in isin)
    total = 0
    for i, c in enumerate(reversed(digits)):
        n = int(c)
        if i % 2 == 1:
            n *= 2
            if n > 9:
                n -= 9
        total += n
    return total % 10 == 0


def _text(value) -> str | None:
    if value is None:
        return None
    text = " ".join(str(value).split())
    return text or None


def read(file: str | Path) -> EuronextList:
    file = Path(file)
    workbook = load_workbook(file, read_only=True, data_only=True)
    sheet = workbook.worksheets[0]
    rows = list(sheet.iter_rows(values_only=True))
    workbook.close()

    list_date = None
    header_index = None
    positions: dict[str, int] = {}
    for i, row in enumerate(rows):
        cells = [_text(v) for v in row]
        if list_date is None:
            for c in cells:
                m = DATE_RE.search(c or "")
                if m:
                    list_date = date(int(m[3]), int(m[2]), int(m[1]))
                    break
        lower = [(c or "").lower() for c in cells]
        if any("isin" in c for c in lower) and any("company" in c for c in lower):
            for key, keyword in COLUMNS.items():
                for j, c in enumerate(lower):
                    if keyword in c:
                        positions[key] = j
                        break
            header_index = i
            break

    if header_index is None or "isin" not in positions or "name" not in positions:
        raise ValueError(f"{file.name}: header row (Company / ISIN) not found")

    companies: list[Company] = []
    rejected: list[tuple[int, str]] = []
    seen: set[str] = set()
    for line_no, row in enumerate(rows[header_index + 1 :], start=header_index + 2):
        def value(key):
            j = positions.get(key)
            return _text(row[j]) if j is not None and j < len(row) else None

        isin, name = (value("isin") or "").upper(), value("name")
        if not isin and not name:
            continue
        if not is_valid_isin(isin):
            rejected.append((line_no, f"invalid ISIN: {isin!r} ({name})"))
            continue
        if isin in seen:
            rejected.append((line_no, f"duplicate ISIN: {isin} ({name})"))
            continue
        seen.add(isin)
        companies.append(Company(isin, name or isin, value("market"), value("segment"), value("country")))

    return EuronextList(file, list_date, companies, rejected)


def compare(conn: sqlite3.Connection, new_list: EuronextList) -> Changes:
    known = {r["isin"]: r for r in conn.execute("SELECT * FROM company")}
    incoming = {c.isin for c in new_list.companies}
    added = [c for c in new_list.companies if c.isin not in known]
    returned = [c for c in new_list.companies if c.isin in known and not known[c.isin]["listed"]]
    removed = [r for isin, r in known.items() if r["listed"] and isin not in incoming]
    unchanged = len(new_list.companies) - len(added) - len(returned)
    return Changes(added, removed, returned, unchanged)


def apply(conn: sqlite3.Connection, new_list: EuronextList, changes: Changes) -> None:
    day = (new_list.list_date or date.today()).isoformat()
    with conn:
        for c in new_list.companies:
            conn.execute(
                """
                INSERT INTO company (isin, name, market, segment, country, listed, added_on, list_date)
                VALUES (?, ?, ?, ?, ?, 1, ?, ?)
                ON CONFLICT(isin) DO UPDATE SET
                    name = excluded.name, market = excluded.market,
                    segment = excluded.segment, country = excluded.country,
                    listed = 1, removed_on = NULL, list_date = excluded.list_date
                """,
                (c.isin, c.name, c.market, c.segment, c.country, day, day),
            )
        for r in changes.removed:
            conn.execute("UPDATE company SET listed = 0, removed_on = ? WHERE isin = ?", (day, r["isin"]))
        conn.execute(
            """
            INSERT INTO list_import (file, list_date, imported_at, company_count, added_count, removed_count)
            VALUES (?, ?, ?, ?, ?, ?)
            """,
            (
                new_list.file.name,
                new_list.list_date.isoformat() if new_list.list_date else None,
                datetime.now().isoformat(timespec="seconds"),
                len(new_list.companies),
                len(changes.added) + len(changes.returned),
                len(changes.removed),
            ),
        )


def latest_list_date(conn: sqlite3.Connection) -> str | None:
    r = conn.execute("SELECT MAX(list_date) AS d FROM list_import").fetchone()
    return r["d"] if r else None
