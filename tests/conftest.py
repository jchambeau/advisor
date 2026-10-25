from pathlib import Path

import pytest
from openpyxl import Workbook

from advisor.db import connect

HEADER = [
    "Société/Company",
    "CodeISIN/ISINCode",
    "Marché/Market",
    "Compartiment/Compartment",
    "Pays d'incorporation/Country of Incorporation",
]


def write_list(path: Path, rows: list[tuple], date_text: str = "23/10/2026 10:00 AM CET") -> Path:
    """Mimic the Euronext file layout: banner, date, disclaimer, header on row 17, columns D to H."""
    wb = Workbook()
    ws = wb.active
    ws.title = "PEA PME"
    ws.cell(row=3, column=7, value="Dernière mise à jour:")
    ws.cell(row=3, column=8, value=date_text)
    ws.cell(row=7, column=4, value="LISTE DES SOCIETES COTEES ELIGIBLE AU DISPOSITIF PEA-PME")
    ws.cell(row=11, column=4, value="La présente liste est donnée à titre exclusivement informatif.")
    for j, value in enumerate(HEADER):
        ws.cell(row=17, column=4 + j, value=value)
    for i, row in enumerate(rows):
        for j, value in enumerate(row):
            ws.cell(row=18 + i, column=4 + j, value=value)
    wb.save(path)
    return path


@pytest.fixture
def conn():
    c = connect(":memory:")
    yield c
    c.close()


@pytest.fixture
def make_list(tmp_path):
    def _make(name, rows, **kw):
        return write_list(tmp_path / name, rows, **kw)
    return _make
